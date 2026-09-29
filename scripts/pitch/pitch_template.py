"""The official Hack2skill template, opened as the deck's base, and the slide surgery it needs.

The template (Google Slides, owner brand@hack2skill.co) is exported as PPTX once and pinned by
SHA-256: if the organisers change it, the build stops until someone looks at the new version. Its
own slides stay in their order and keep their frame (header band, section title, foot bar); the
deck fills them, removes the instruction page and the two optional sections it does not use, and
clones a section's slide when that section needs a second page.
"""

from __future__ import annotations

import copy
import hashlib
import urllib.request

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from pitch_paths import CACHE, TEMPLATE, TEMPLATE_URL

TEMPLATE_SHA256 = "85c18bf09c9ccf96b51ff4ac82b4ba26fd3bebfbb89fe05940b267bef5665237"

# The template's pages, 1-based, as the organisers ordered them.
COVER = 1
INSTRUCTIONS = 2
BRIEF = 3
OPPORTUNITIES = 4
FEATURES = 5
PROCESS_FLOW = 6
WIREFRAMES = 7  # optional
ARCHITECTURE = 8
TECHNOLOGIES = 9
COST = 10  # optional
SNAPSHOTS = 11
PERFORMANCE = 12
ADDITIONAL = 13
LINKS = 14
SPARE = 15
THANK_YOU = 16
TEMPLATE_PAGES = 16

EXPECTED_TITLES = {
    BRIEF: "Brief about the idea",
    OPPORTUNITIES: "Opportunities",
    FEATURES: "List of features offered by the solution",
    PROCESS_FLOW: "Process flow diagram or Use-case diagram",
    WIREFRAMES: "Wireframes/Mock diagrams of the proposed solution (optional)",
    ARCHITECTURE: "Architecture diagram of the proposed solution",
    TECHNOLOGIES: "Technologies to be used in the solution",
    COST: "Estimated implementation cost (optional)",
    SNAPSHOTS: "Snapshots of the prototype",
    PERFORMANCE: "Prototype Performance report/Benchmarking",
    ADDITIONAL: "Additional Details/Future Development (if any)",
    LINKS: "Provide links to your:",
}


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def open_template() -> Presentation:
    """The pinned template; downloads it once, stops if its bytes changed."""
    if not TEMPLATE.exists():
        CACHE.mkdir(parents=True, exist_ok=True)
        urllib.request.urlretrieve(TEMPLATE_URL, TEMPLATE)
    actual = _sha256(TEMPLATE.read_bytes())
    if actual != TEMPLATE_SHA256:
        raise RuntimeError(
            f"{TEMPLATE} has SHA-256 {actual}, expected {TEMPLATE_SHA256}: the template changed; "
            "review the new version before pinning it"
        )
    prs = Presentation(TEMPLATE)
    if len(prs.slides) != TEMPLATE_PAGES:
        raise RuntimeError(f"the template has {len(prs.slides)} pages, expected {TEMPLATE_PAGES}")
    for page, expected in EXPECTED_TITLES.items():
        title = title_box(prs.slides[page - 1]).text_frame.paragraphs[0].text.strip()
        if title != expected:
            raise RuntimeError(f"template page {page} is titled {title!r}, expected {expected!r}")
    return prs


def title_box(slide: object) -> object:
    """The section title: the one text box of a template content page."""
    boxes = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]  # type: ignore[attr-defined]
    if len(boxes) != 1:
        raise RuntimeError(f"expected one titled text box, found {len(boxes)}")
    return boxes[0]


def set_title(slide: object, title: str) -> None:
    """Keeps the title's own run format; drops the template's prompt lines under it."""
    frame = title_box(slide).text_frame
    first = frame.paragraphs[0]
    runs = first.runs
    runs[0].text = title
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    for br in first._p.findall(qn("a:br")):
        first._p.remove(br)
    for extra in frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def clone_slide(prs: Presentation, source: object) -> object:
    """A copy of `source` (frame, title and all) appended at the end; move it with move_slide."""
    copy_ = prs.slides.add_slide(source.slide_layout)  # type: ignore[attr-defined]
    for shape in list(copy_.shapes):
        shape._element.getparent().remove(shape._element)
    tree = copy_.shapes._spTree
    for shape in source.shapes:  # type: ignore[attr-defined]
        element = copy.deepcopy(shape._element)
        for blip in element.iter(qn("a:blip")):
            image_part = source.part.related_part(blip.get(qn("r:embed")))  # type: ignore[attr-defined]
            blip.set(qn("r:embed"), copy_.part.relate_to(image_part, RT.IMAGE))
        tree.append(element)
    return copy_


def _slide_ids(prs: Presentation) -> list[object]:
    return list(prs.slides._sldIdLst)  # noqa: SLF001 — python-pptx has no public reorder API


def move_slide(prs: Presentation, slide: object, index: int) -> None:
    """Moves `slide` to 0-based `index` in the deck's order."""
    ids = prs.slides._sldIdLst  # noqa: SLF001
    for sld_id in _slide_ids(prs):
        if prs.part.related_part(sld_id.rId) is slide.part:  # type: ignore[attr-defined]
            ids.remove(sld_id)
            ids.insert(index, sld_id)
            return
    raise RuntimeError("slide not found in the deck")


def delete_slide(prs: Presentation, slide: object) -> None:
    ids = prs.slides._sldIdLst  # noqa: SLF001
    for sld_id in _slide_ids(prs):
        if prs.part.related_part(sld_id.rId) is slide.part:  # type: ignore[attr-defined]
            prs.part.drop_rel(sld_id.rId)
            ids.remove(sld_id)
            return
    raise RuntimeError("slide not found in the deck")


def index_of(prs: Presentation, slide: object) -> int:
    for i, s in enumerate(prs.slides):
        if s.part is slide.part:  # type: ignore[attr-defined]
            return i
    raise RuntimeError("slide not found in the deck")
