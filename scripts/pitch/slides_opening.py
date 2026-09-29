"""The opening pages: the template's cover (team details), the idea, two opportunity pages."""

from __future__ import annotations

from pptx.util import Emu, Pt

from pitch_draw import Canvas, box, p, r, text
from pitch_fonts import FAMILY_MEDIUM, FAMILY_SEMIBOLD
from pitch_notes import NoteBook
from pitch_parts import credit, headline
from pitch_paths import STILLS
from pitch_shapes import line, picture, rect
from pitch_theme import (
    AMBER,
    AMBER_EDGE,
    BODY,
    BODY_STRONG,
    FINE,
    INK,
    LABEL,
    LEAD,
    RULE_STRONG,
    SMALL,
    SPEECH,
)

COVER_VALUE_PT = 18
COVER_TEXT_W = Emu(int(5.9 * 914400))
# Labels whose value starts on its own line, under the label's text: the theme's name is too long
# to follow "Problem Statement:" on one line at the template's 18 pt (owner, 2026-09-29).
VALUE_BELOW = {"Problem Statement:"}
MONTHS = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()


def fill_cover(slide: object, data: dict, canvas: Canvas) -> None:
    """Writes the team's details after the template's own labels and adds the product's name."""
    details = next(
        s for s in slide.shapes if s.has_text_frame and "Team Details" in s.text_frame.text
    )  # type: ignore[attr-defined]
    details.width = COVER_TEXT_W
    team = data["team"]
    values = {
        "Team name:": team["name"],
        "Team leader name:": team["leader"] or "",
        # The chosen theme: the FAQ reads "one solution under one problem statement/theme".
        "Problem Statement:": team["theme"],
    }
    filled = 0
    for paragraph in details.text_frame.paragraphs:
        label = paragraph.text.strip()
        if label not in values:
            continue
        if label in VALUE_BELOW:
            paragraph.add_line_break()
        run = paragraph.add_run()
        run.text = values[label]
        run.font.name = FAMILY_MEDIUM
        run.font.size = Pt(COVER_VALUE_PT)
        run.font.color.rgb = INK
        filled += 1
    if filled != len(values):
        raise RuntimeError(f"the cover's team details had {filled} of {len(values)} labels")
    text(
        canvas,
        box(6.35, 3.28, 3.3, 1.72),
        [
            p(r("Gapline", (FAMILY_SEMIBOLD, 34, INK))),
            p(r("Descriptions that fit", LEAD)),
            p(r("between the lines.", LEAD), space_after=8),
            # The themes page also asks for a category (Healthcare, Education, ... Accessibility).
            p(r(f"Category: {team['category']}", LABEL)),
        ],
        "product",
    )


def _seven_exhibit(canvas: Canvas, data: dict, left: float, top: float, width: float) -> float:
    """Two stills, the silence between two lines of dialogue, and the lines Gapline fit in it."""
    seven = data["seven"]
    gap = 0.08
    still_w = (width - gap) / 2
    still_h = still_w / 2.4
    for i, name in enumerate(("seven-1.jpg", "seven-2.jpg")):
        picture(canvas, STILLS / name, box(left + i * (still_w + gap), top, still_w, still_h))
    y = top + still_h + 0.08
    for i, cue in enumerate(seven["lines"]):
        x = left + i * (still_w + gap)
        rect(canvas, box(x, y + 0.06, 0.09, 0.09), fill=AMBER, line=AMBER_EDGE)
        text(
            canvas,
            box(x + 0.15, y, still_w - 0.15, 0.42),
            [p(r(cue["text"], SMALL))],
            f"line-{cue['id']}",
        )

    t0, t1 = seven["domain"]

    def at(t: float) -> float:
        return left + (t - t0) / (t1 - t0) * width

    def emu(inches: float) -> Emu:
        return Emu(int(inches * 914400))

    lane = y + 0.46 + 0.34
    lane_h = 0.24
    x0, x1 = at(seven["locked"]["end"]), at(seven["freaky"]["start"])
    bracket = lane - 0.07
    line(canvas, emu(x0), emu(bracket), emu(x1), emu(bracket), RULE_STRONG)
    for x in (x0, x1):
        line(canvas, emu(x), emu(bracket), emu(x), emu(bracket + 0.05), RULE_STRONG)
    text(
        canvas,
        box(x0, bracket - 0.26, x1 - x0, 0.22),
        [p(r(f"{seven['silence']:.2f} s with no dialogue", SMALL), align="c")],
        "silence",
    )
    for quote, span, align in (
        ("“…locked.”", seven["locked"], "l"),
        ("“This is pretty freaky.”", seven["freaky"], "r"),
    ):
        rect(
            canvas,
            box(at(span["start"]), lane, at(span["end"]) - at(span["start"]), lane_h),
            fill=SPEECH,
        )
        label_x = at(span["start"]) if align == "l" else at(span["end"]) - 2.0
        text(
            canvas,
            box(label_x, lane + lane_h + 0.04, 2.0, 0.2),
            [p(r(quote, FINE), align=align)],
            f"quote-{align}",
        )
    for cue in seven["lines"]:
        start = cue["start"] + cue["onset"]
        end = start + cue["voiced"]
        where = box(at(start), lane, at(end) - at(start), lane_h)
        rect(canvas, where, fill=AMBER, line=AMBER_EDGE)
        text(
            canvas,
            where,
            [p(r(f"{cue['voiced']:.2f} s", (FINE[0], FINE[1], INK)), align="c")],
            f"bar-{cue['id']}",
            anchor="m",
        )
    return lane + lane_h + 0.26


def build_brief(canvas: Canvas, data: dict, notes: NoteBook) -> None:
    top = headline(canvas, "Gapline writes audio description that fits between the lines.")
    section = "Brief about the idea"
    left_w = 3.55
    text(
        canvas,
        box(0.39, top, left_w, 5.1 - top),
        [
            p(
                r(
                    "Audio description tells blind and low-vision viewers what is on screen, in "
                    "the pauses between the dialogue. Made by hand, one film takes about three "
                    "months.",
                    BODY_STRONG,
                ),
                r(notes.mark("hand_made", section), BODY_STRONG, sup=True),
                space_after=10,
            ),
            p(
                r(
                    "Upload a clip and press Generate. Gapline maps the silences, writes a line "
                    "for each with Gemini, reviews it against published guidelines, voices and "
                    "measures it, and mixes the described film, in Korean or English.",
                    BODY,
                ),
            ),
        ],
        "idea",
    )
    exhibit_left = 0.39 + left_w + 0.35
    bottom = _seven_exhibit(canvas, data, exhibit_left, top, 9.61 - exhibit_left)
    text(
        canvas,
        box(exhibit_left, bottom + 0.1, 9.61 - exhibit_left, 0.42),
        [
            p(
                r(
                    "Tears of Steel, 53.7–60.9 s: a blind viewer hears only a hum. Gapline fit two "
                    "lines into the silence; each ends before the next word.",
                    SMALL,
                )
            )
        ],
        "seven-caption",
    )
    credit(canvas, data["filmCredit"])
