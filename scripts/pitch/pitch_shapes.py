"""Shapes for the template deck: rectangles, lines, arrows and pictures, free of theme styles.

Theme styles are stripped from every shape so no renderer adds shadows or theme fills.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn

from pitch_draw import Box, Canvas
from pitch_theme import HAIRLINE


def _strip_style(shape: object) -> None:
    style = shape._element.find(qn("p:style"))  # type: ignore[attr-defined]
    if style is not None:
        shape._element.remove(style)  # type: ignore[attr-defined]


def rect(
    canvas: Canvas,
    where: Box,
    fill: object | None = None,
    line: object | None = None,
    line_w: object = HAIRLINE,
    radius: float | None = None,
    dash: bool = False,
    name: str | None = None,
) -> object:
    """A filled or outlined rectangle; `radius` in inches rounds its corners."""
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = canvas.slide.shapes.add_shape(kind, where.x, where.y, where.w, where.h)  # type: ignore[attr-defined]
    _strip_style(shape)
    if name:
        shape.name = name
    if radius:
        shape.adjustments[0] = min(0.5, radius * 914400 / min(where.w, where.h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = line_w
        if dash:
            ln = shape.line._get_or_add_ln()
            dash_el = etree.SubElement(ln, qn("a:prstDash"))
            dash_el.set("val", "dash")
    return shape


def line(
    canvas: Canvas,
    x1: int,
    y1: int,
    x2: int,
    y2: int,
    color: object,
    width: object = HAIRLINE,
    arrow: bool = False,
    dash: bool = False,
    both: bool = False,
) -> object:
    """A straight line; `arrow` puts a head at (x2, y2), `both` at each end."""
    shape = canvas.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)  # type: ignore[attr-defined]
    _strip_style(shape)
    shape.line.color.rgb = color
    shape.line.width = width
    ln = shape.line._get_or_add_ln()
    if dash:
        dash_el = etree.SubElement(ln, qn("a:prstDash"))
        dash_el.set("val", "dash")
    heads = (("a:headEnd",) if both else ()) + (("a:tailEnd",) if arrow or both else ())
    for tag in heads:
        head = etree.SubElement(ln, qn(tag))
        head.set("type", "triangle")
        head.set("w", "med")
        head.set("len", "med")
    return shape


def arrow(
    canvas: Canvas, points: list[tuple[float, float]], color: object, both: bool = False
) -> None:
    """A path of straight segments through `points` (inches); the last segment ends in a head."""
    emu = [(int(x * 914400), int(y * 914400)) for x, y in points]
    last = len(emu) - 2
    for i, ((x1, y1), (x2, y2)) in enumerate(zip(emu, emu[1:])):
        line(canvas, x1, y1, x2, y2, color, arrow=i == last, both=both and i == 0 and last == 0)


def picture(
    canvas: Canvas, path: Path, where: Box, border: object | None = None, stretch: bool = False
) -> object:
    """A picture filling `where`, cropped to its aspect (centred); `stretch` keeps it whole."""
    with Image.open(path) as im:
        iw, ih = im.size
    shape = canvas.slide.shapes.add_picture(str(path), where.x, where.y, where.w, where.h)  # type: ignore[attr-defined]
    if stretch:
        return shape
    target = where.w / where.h
    source = iw / ih
    if source > target:
        cut = (1 - target / source) / 2
        shape.crop_left = shape.crop_right = cut
    elif source < target:
        cut = (1 - source / target) / 2
        shape.crop_top = shape.crop_bottom = cut
    if border is not None:
        shape.line.color.rgb = border
        shape.line.width = HAIRLINE
    return shape
